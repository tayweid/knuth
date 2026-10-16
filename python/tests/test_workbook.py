"""Excel workbooks: one sheet in, CSV text out (workbook.py), opened by path
and never written (files.py)."""

import csv
import datetime
import io

from openpyxl import Workbook

from knuth import files, workbook
from knuth.server import _validate_request


def _gradebook(path):
    """A small workbook the way a teacher's looks: a header row, mixed
    values, a formula, a second sheet, and a hidden one."""
    book = Workbook()
    grades = book.active
    grades.title = "Grades"
    grades.append(["Name", "Score", "Passed", "Due", "Note"])
    grades.append(["Ada", 93.5, True, datetime.datetime(2026, 10, 16), None])
    grades.append(["Bob, Jr.", 0.1 + 0.2, False, datetime.datetime(2026, 10, 16, 9, 30), 'said "hi"'])
    grades.append(["Cy\nTwo lines", 7, None, datetime.date(2026, 1, 2)])
    grades["A5"] = "Total"
    grades["B5"] = "=SUM(B2:B4)"  # no cached value: never computed by Excel
    grades["A7"] = None  # trailing empty rows are dropped
    notes = book.create_sheet("Notes")
    notes["B2"] = "only cell"
    hidden = book.create_sheet("Answers")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "secret"
    book.save(path)
    return path


def test_a_sheet_reads_as_a_rectangle_of_text(tmp_path):
    path = _gradebook(tmp_path / "grades.xlsx")
    text, book = workbook.read_sheet(path)
    assert book == {"sheets": ["Grades", "Notes"], "sheet": "Grades", "rows": 5, "truncated": False}
    rows = list(csv.reader(io.StringIO(text)))
    assert rows == [
        ["Name", "Score", "Passed", "Due", "Note"],
        ["Ada", "93.5", "TRUE", "2026-10-16", ""],
        ["Bob, Jr.", "0.3", "FALSE", "2026-10-16 09:30:00", 'said "hi"'],
        ["Cy\nTwo lines", "7", "", "2026-01-02", ""],
        ["Total", "", "", "", ""],
    ]
    # Quoted only where CSV must be, so the grid reads it back unchanged.
    assert '"Bob, Jr."' in text and '"said ""hi"""' in text and text.endswith(",,,,\n")


def test_another_sheet_by_name_and_the_first_when_it_is_gone(tmp_path):
    path = _gradebook(tmp_path / "grades.xlsx")
    text, book = workbook.read_sheet(path, "Notes")
    assert book["sheet"] == "Notes"
    # Empty leading rows and columns stay: B2 is still B2.
    assert text == ",\n,only cell\n"
    assert workbook.read_sheet(path, "Renamed")[1]["sheet"] == "Grades"
    # A hidden sheet is not offered, even by name.
    assert workbook.read_sheet(path, "Answers")[1]["sheet"] == "Grades"


def test_a_sheet_past_the_limit_arrives_cut_short(tmp_path):
    path = tmp_path / "big.xlsx"
    book = Workbook()
    for n in range(200):
        book.active.append([n, "x" * 20])
    book.save(path)
    text, info = workbook.read_sheet(path, limit=1000)
    assert info["truncated"] is True
    assert 0 < info["rows"] < 200 and len(text.encode()) <= 1000
    assert text.splitlines()[0] == "0," + "x" * 20


def test_an_empty_sheet_is_empty_text(tmp_path):
    path = tmp_path / "empty.xlsx"
    Workbook().save(path)
    assert workbook.read_sheet(path) == ("", {"sheets": ["Sheet"], "sheet": "Sheet", "rows": 0, "truncated": False})


def test_open_document_shows_a_workbook_and_save_never_writes_one(tmp_path):
    path = _gradebook(tmp_path / "grades.xlsx")
    before = path.read_bytes()
    reply = files.open_document(str(path), "Notes")
    assert reply["path"] == str(path) and reply["name"] == "grades.xlsx"
    assert reply["text"] == ",\n,only cell\n"
    assert reply["modified"] == int(path.stat().st_mtime * 1000)
    assert reply["workbook"]["sheet"] == "Notes" and "unsaved" not in reply
    assert "only reads" in files.save_document(str(path), "a,b\n")["error"]
    assert "only reads" in files.save_document(str(tmp_path / "new.XLSM"), "a,b\n")["error"]
    assert path.read_bytes() == before
    # The sheet saved as a CSV is the text as given: no Python header.
    saved = files.save_document(str(tmp_path / "grades.csv"), reply["text"])
    assert "header" not in saved
    assert (tmp_path / "grades.csv").read_text() == ",\n,only cell\n"


def test_open_names_a_file_that_is_not_a_workbook(tmp_path):
    path = tmp_path / "fake.xlsx"
    path.write_bytes(b"not a zip at all")
    assert files.open_document(str(path)) == {"error": "fake.xlsx could not be read as an Excel workbook"}


def test_the_open_request_takes_a_sheet_name():
    assert _validate_request({"type": "open", "id": 1, "path": "/p/g.xlsx", "sheet": "Notes"}) is None
    assert _validate_request({"type": "open", "id": 1, "path": "/p/g.xlsx"}) is None
    assert "sheet" in _validate_request({"type": "open", "id": 1, "path": "/p/g.xlsx", "sheet": 3})
    assert "sheet" in _validate_request({"type": "open", "id": 1, "path": "/p/g.xlsx", "sheet": "x" * 300})
