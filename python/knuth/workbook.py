"""Read Excel workbooks: one sheet of an .xlsx in, CSV text out.

Read-only by design. The page shows the sheet as its grid of cells (the
same grid a .csv opens in, src/grid.ts) with nothing to edit, and saving
writes a CSV elsewhere — never the workbook, which files.py refuses to
write. Implemented once, here, the way the notebook import is (ipynb.py).

What a cell shows (decided 2026-10-16):
- the value Excel last computed (`data_only`): a formula shows its cached
  result, and nothing in a file no spreadsheet app has saved since it was
  written;
- the value, not its number format: 0.5 where Excel shows 50%, and a
  number to Excel's own 15 significant digits;
- a date as ISO (2026-10-16), with its time only when it has one;
- TRUE and FALSE as Excel writes them.
Hidden sheets are left out, as Excel leaves them out of its tabs. A sheet
larger than MAX_DOCUMENT_BYTES as text arrives cut short, and says so.
"""

import csv
import datetime
import io
import warnings

from .limits import MAX_DOCUMENT_BYTES

SUFFIXES = (".xlsx", ".xlsm")


def is_workbook(path):
    return path.suffix.lower() in SUFFIXES


def _text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return format(value, ".15g")
    if isinstance(value, datetime.datetime):
        if value.time() == datetime.time(0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    return str(value)


def read_sheet(path, sheet=None, limit=MAX_DOCUMENT_BYTES):
    """One worksheet of the workbook at `path` as CSV text: the sheet named
    `sheet`, or the first visible one when that is None or no longer there.

    Returns (text, workbook) where workbook is {sheets, sheet, rows,
    truncated}: the visible sheets' names in order, the one shown, how many
    rows the text holds, and whether the sheet had more. Every row is as
    wide as the widest, so the text reads back as a table anywhere; trailing
    empty rows and columns are dropped. Raises ValueError with a sentence
    for the page.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError(
            f"{path.name} needs openpyxl to open, and Knuth's engine does not have it"
            " (python3 -m pip install openpyxl)"
        ) from None
    with warnings.catch_warnings():
        # openpyxl warns about parts it drops (data validation, conditional
        # formatting); a reader that keeps only values loses nothing.
        warnings.simplefilter("ignore")
        try:
            book = load_workbook(path, read_only=True, data_only=True)
        except Exception:  # a damaged or foreign zip fails in many ways
            raise ValueError(f"{path.name} could not be read as an Excel workbook") from None
        rows = []
        truncated = False
        try:
            visible = [ws for ws in book.worksheets if ws.sheet_state == "visible"]
            ws = next((each for each in visible if each.title == sheet), visible[0] if visible else None)
            if ws is not None:
                # The size a workbook declares can be wrong (or the whole
                # sheet); the rows themselves are what is there.
                ws.reset_dimensions()
                size = 0
                for values in ws.iter_rows(values_only=True):
                    cells = [_text(value) for value in values]
                    while cells and cells[-1] == "":
                        cells.pop()
                    size += sum(len(cell) for cell in cells) + len(cells) + 1
                    if size > limit:
                        truncated = True
                        break
                    rows.append(cells)
        except Exception:
            raise ValueError(f"{path.name} could not be read as an Excel workbook") from None
        finally:
            book.close()
    if ws is None:
        raise ValueError(f"{path.name} has no sheets of cells")
    while rows and not rows[-1]:
        rows.pop()
    width = max(map(len, rows), default=0)
    lines = []
    total = 0
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    for cells in rows:
        buffer.seek(0)
        buffer.truncate()
        writer.writerow(cells + [""] * (width - len(cells)))
        line = buffer.getvalue()
        total += len(line.encode("utf-8"))
        if total > limit:
            truncated = True
            break
        lines.append(line)
    return "".join(lines), {
        "sheets": [each.title for each in visible],
        "sheet": ws.title,
        "rows": len(lines),
        "truncated": truncated,
    }
